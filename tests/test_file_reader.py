from io import BytesIO

import pandas as pd
from openpyxl import Workbook

from modules.file_reader import read_uploaded_file


def test_csv_preserves_leading_zero_identifier():
    data = b"student_id,postal_code\n00125,00501\n00126,00502\n"
    datasets = read_uploaded_file("people.csv", data)
    assert len(datasets) == 1
    assert datasets[0].frame.loc[0, "student_id"] == "00125"
    assert datasets[0].frame.loc[0, "postal_code"] == "00501"


def test_csv_duplicate_headers_are_renamed_and_reported():
    datasets = read_uploaded_file("dupe.csv", b"id,id,\n1,2,3\n")
    assert datasets[0].frame.columns.tolist() == ["id", "id_2", "column_3"]
    assert any("Duplicate column" in warning for warning in datasets[0].warnings)


def test_excel_reads_all_worksheets():
    workbook = Workbook()
    ws = workbook.active
    ws.title = "Students"
    ws.append(["student_id", "name"])
    ws.append(["S001", "Ada"])
    ws2 = workbook.create_sheet("Courses")
    ws2.append(["course_id", "title"])
    ws2.append(["C001", "SQL"])
    buffer = BytesIO()
    workbook.save(buffer)
    datasets = read_uploaded_file("school.xlsx", buffer.getvalue())
    assert {dataset.source_sheet for dataset in datasets} == {"Students", "Courses"}
    assert all(len(dataset.frame) == 1 for dataset in datasets)


def test_json_records_parse():
    datasets = read_uploaded_file("records.json", b'[{"id":"001","meta":{"city":"Delhi"}},{"id":"002","meta":{"city":"Noida"}}]')
    assert len(datasets) == 1
    assert datasets[0].frame.columns.tolist() == ["id", "meta__city"]
    assert datasets[0].frame.loc[0, "id"] == "001"


def test_invalid_json_reports_actionable_error():
    try:
        read_uploaded_file("broken.json", b"{broken")
    except ValueError as exc:
        assert "Invalid JSON" in str(exc)
    else:
        raise AssertionError("Invalid JSON should not parse")


def test_headerless_csv_preserves_first_row_and_assigns_generic_columns():
    datasets = read_uploaded_file("no_header.csv", b"00125,Asha,Delhi\n00126,Kabir,Noida\n", has_header=False)
    frame = datasets[0].frame
    assert frame.columns.tolist() == ["column_1", "column_2", "column_3"]
    assert len(frame) == 2
    assert frame.loc[0, "column_1"] == "00125"
    assert datasets[0].source_record_offset == 1
    assert any("all source rows were preserved" in warning for warning in datasets[0].warnings)


def test_json_nested_lists_are_serialised_safely():
    datasets = read_uploaded_file("nested.json", b'[{"id":"1","tags":["sql","python"]}]')
    value = datasets[0].frame.loc[0, "tags"]
    assert isinstance(value, str)
    assert "sql" in value


def test_xlsm_reads_values_without_executing_macros():
    workbook = Workbook()
    ws = workbook.active
    ws.title = "Data"
    ws.append(["id", "name"])
    ws.append(["001", "Ada"])
    buffer = BytesIO()
    workbook.save(buffer)
    datasets = read_uploaded_file("data.xlsm", buffer.getvalue())
    assert datasets[0].frame.loc[0, "id"] == "001"
    assert datasets[0].source_sheet == "Data"


def test_xml_repeating_rows_parse_as_tabular_data():
    xml = b"<data><row><id>1</id><name>Ada</name></row><row><id>2</id><name>Lin</name></row></data>"
    datasets = read_uploaded_file("people.xml", xml)
    assert datasets[0].frame.columns.tolist() == ["id", "name"]
    assert len(datasets[0].frame) == 2


def test_zip_archive_parses_supported_csv_member():
    import zipfile
    archive_buffer = BytesIO()
    with zipfile.ZipFile(archive_buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("data/students.csv", "student_id,name\nS001,Ada\n")
    datasets = read_uploaded_file("bundle.zip", archive_buffer.getvalue())
    assert len(datasets) == 1
    assert datasets[0].source_sheet == "data/students.csv"
    assert datasets[0].frame.loc[0, "student_id"] == "S001"


def test_legacy_xls_reading_when_test_writer_and_reader_are_installed():
    import pytest
    xlwt = pytest.importorskip("xlwt")
    pytest.importorskip("xlrd")
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("People")
    for col, value in enumerate(["id", "name"]):
        sheet.write(0, col, value)
    sheet.write(1, 0, "001")
    sheet.write(1, 1, "Ada")
    buffer = BytesIO()
    workbook.save(buffer)
    datasets = read_uploaded_file("people.xls", buffer.getvalue())
    assert datasets[0].frame.loc[0, "id"] == "001"


def test_parquet_reading_when_pyarrow_is_installed():
    import pytest
    pyarrow = pytest.importorskip("pyarrow")
    frame = pd.DataFrame({"id": ["001", "002"], "name": ["Ada", "Lin"]})
    buffer = BytesIO()
    frame.to_parquet(buffer, engine="pyarrow")
    datasets = read_uploaded_file("people.parquet", buffer.getvalue())
    assert datasets[0].frame.loc[0, "id"] == "001"


def test_excel_duplicate_headers_are_not_silently_lost():
    workbook = Workbook()
    ws = workbook.active
    ws.append(["id", "id"])
    ws.append(["001", "002"])
    buffer = BytesIO()
    workbook.save(buffer)
    datasets = read_uploaded_file("dupe.xlsx", buffer.getvalue())
    assert datasets[0].frame.columns.tolist() == ["id", "id_2"]
    assert datasets[0].frame.loc[0, "id_2"] == "002"


def test_empty_excel_worksheet_is_preserved_with_warning():
    workbook = Workbook()
    workbook.active.title = "Empty"
    data = workbook.create_sheet("Data")
    data.append(["id", "name"])
    data.append(["S1", "Ada"])
    buffer = BytesIO()
    workbook.save(buffer)
    datasets = read_uploaded_file("with_empty.xlsx", buffer.getvalue())
    empty = next(dataset for dataset in datasets if dataset.source_sheet == "Empty")
    assert empty.frame.empty
    assert any("empty" in warning.lower() for warning in empty.warnings)
